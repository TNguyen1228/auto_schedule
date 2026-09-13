import pandas as pd
from typing import List, Dict, Optional
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from scheduler.models import (
    ScheduleConfig,
    ClassGroup,
    Teacher,
    TeachingAssignment,
    ScheduleResult
)

class TimetableExporter:
    """
    Xuất kết quả Thời Khóa Biểu ra Excel và Terminal:
    - Hiển thị rõ 2 buổi: Sáng (Tiết 1-4) & Chiều (Tiết 5-7)
    - Đánh dấu [NGHỈ CHIỀU] vào chiều Thứ 6
    """
    def __init__(
        self,
        config: ScheduleConfig,
        classes: List[ClassGroup],
        teachers: List[Teacher],
        assignments: List[TeachingAssignment],
        result: ScheduleResult
    ):
        self.config = config
        self.classes = classes
        self.teachers = teachers
        self.assignments = assignments
        self.result = result

    def print_summary(self):
        """In tóm tắt kết quả ra màn hình Console"""
        print("=" * 70)
        print(f"TRẠNG THÁI: {self.result.status}")
        print(f"THỜI GIAN GIẢI: {self.result.solve_time_seconds:.3f} giây")
        print(f"TỔNG SỐ TIẾT ĐÃ XẾP: {len(self.result.slots)}")
        print(f"KHUNG HỌC: Sáng 4 tiết, Chiều 3 tiết, Chiều Thứ 6 nghỉ ({self.config.total_active_slots_per_week} tiết/tuần)")
        print(f"THÔNG ĐIỆP: {self.result.message}")
        print("=" * 70)

    # def print_class_timetable(self, class_id: str):
    #     """In TKB của một lớp ra console dạng bảng phân tách Sáng/Chiều"""
    #     class_slots = [s for s in self.result.slots if s.class_id == class_id]
    #     grid = {d: {p: "" for p in self.config.all_periods} for d in self.config.days}
    #     for s in class_slots:
    #         grid[s.day][s.period] = f"{s.subject} ({s.teacher_name})"

    #     # Gán ô nghỉ chiều Thứ 6
    #     for (d, p) in self.config.closed_slots:
    #         grid[d][p] = "[NGHỈ CHIỀU]"

    #     print(f"\n--- THỜI KHÓA BIỂU LỚP {class_id} (TIỂU HỌC 2 BUỔI/NGÀY) ---")
    #     header = f"{'Buổi':<7} | {'Tiết':<6} | " + " | ".join(f"{d:<24}" for d in self.config.days)
    #     print("-" * len(header))
    #     print(header)
    #     print("-" * len(header))

    #     for session_name, periods in [("SÁNG", self.config.morning_periods), ("CHIỀU", self.config.afternoon_periods)]:
    #         for p in periods:
    #             s_label = session_name if p == periods[0] else ""
    #             row_str = f"{s_label:<7} | Tiết {p:<2} | "
    #             cols = []
    #             for d in self.config.days:
    #                 val = grid[d][p]
    #                 cols.append(f"{val:<24}")
    #             row_str += " | ".join(cols)
    #             print(row_str)
    #         print("-" * len(header))

    def export_to_excel(self, file_path: str = "Thoi_Khoa_Bieu_Truong_9_Lop.xlsx"):
        """
        Xuất file Excel chuyên nghiệp với phân chia rõ ràng Sáng (1-4) & Chiều (5-7)
        """
        class_slot_map = {(s.class_id, s.day, s.period): s for s in self.result.slots}
        teacher_slot_map = {}
        for s in self.result.slots:
            key = (s.teacher_id, s.day, s.period)
            if key not in teacher_slot_map:
                teacher_slot_map[key] = []
            teacher_slot_map[key].append(s)

        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        session_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        closed_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='BFBFBF'),
            right=Side(style='thin', color='BFBFBF'),
            top=Side(style='thin', color='BFBFBF'),
            bottom=Side(style='thin', color='BFBFBF')
        )

        # -------------------------------------------------------------
        # SHEET 1: THỜI KHÓA BIỂU THEO LỚP
        # -------------------------------------------------------------
        ws_class = wb.create_sheet(title="TKB_Theo_Lop")
        ws_class.views.sheetView[0].showGridLines = True

        ws_class.cell(row=1, column=1, value="BẢNG THỜI KHÓA BIỂU THEO LỚP HỌC (TIỂU HỌC)").font = Font(size=14, bold=True, color="1F4E79")
        ws_class.cell(row=2, column=1, value="Sáng 4 tiết, Chiều 3 tiết - Chiều Thứ 6 nghỉ").font = Font(size=10, italic=True)

        current_row = 4
        ws_class.cell(row=current_row, column=1, value="Ngày").fill = header_fill
        ws_class.cell(row=current_row, column=1).font = header_font
        ws_class.cell(row=current_row, column=1).alignment = Alignment(horizontal='center', vertical='center')

        ws_class.cell(row=current_row, column=2, value="Buổi").fill = header_fill
        ws_class.cell(row=current_row, column=2).font = header_font
        ws_class.cell(row=current_row, column=2).alignment = Alignment(horizontal='center', vertical='center')

        ws_class.cell(row=current_row, column=3, value="Tiết").fill = header_fill
        ws_class.cell(row=current_row, column=3).font = header_font
        ws_class.cell(row=current_row, column=3).alignment = Alignment(horizontal='center', vertical='center')

        col_idx = 4
        for c in self.classes:
            cell = ws_class.cell(row=current_row, column=col_idx, value=f"{c.id}")
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            col_idx += 1

        current_row += 1
        day_fill_toggle = [PatternFill(start_color="FFFFFF", fill_type="solid"), PatternFill(start_color="F9FBFD", fill_type="solid")]
        
        day_idx = 0
        for d in self.config.days:
            start_row_for_day = current_row
            row_fill = day_fill_toggle[day_idx % 2]
            day_idx += 1

            for session_name, periods in [("Sáng", self.config.morning_periods), ("Chiều", self.config.afternoon_periods)]:
                start_row_session = current_row
                for p in periods:
                    ws_class.cell(row=current_row, column=3, value=f"Tiết {p}").alignment = Alignment(horizontal='center', vertical='center')
                    ws_class.cell(row=current_row, column=3).border = thin_border
                    ws_class.cell(row=current_row, column=3).fill = row_fill

                    col = 4
                    is_closed = (d, p) in self.config.closed_slots
                    for c in self.classes:
                        cell = ws_class.cell(row=current_row, column=col)
                        cell.border = thin_border
                        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

                        if is_closed:
                            cell.value = "[NGHỈ CHIỀU]"
                            cell.fill = closed_fill
                            cell.font = Font(color="7F7F7F", italic=True)
                        else:
                            slot = class_slot_map.get((c.id, d, p))
                            cell.fill = row_fill
                            if slot:
                                cell.value = f"{slot.subject}\n({slot.teacher_name})"
                                if "Chào cờ" in slot.subject:
                                    cell.fill = PatternFill(start_color="FFF2CC", fill_type="solid")
                                elif "Sinh hoạt" in slot.subject:
                                    cell.fill = PatternFill(start_color="E2EFDA", fill_type="solid")
                            else:
                                cell.value = ""
                        col += 1
                    current_row += 1

                # Merge ô Buổi
                ws_class.merge_cells(start_row=start_row_session, start_column=2, end_row=current_row - 1, end_column=2)
                sess_cell = ws_class.cell(row=start_row_session, column=2, value=session_name)
                sess_cell.alignment = Alignment(horizontal='center', vertical='center')
                sess_cell.font = Font(bold=True)
                sess_cell.fill = session_fill
                for r in range(start_row_session, current_row):
                    ws_class.cell(row=r, column=2).border = thin_border

            # Merge ô Ngày
            ws_class.merge_cells(start_row=start_row_for_day, start_column=1, end_row=current_row - 1, end_column=1)
            day_cell = ws_class.cell(row=start_row_for_day, column=1, value=d)
            day_cell.alignment = Alignment(horizontal='center', vertical='center')
            day_cell.font = Font(bold=True)
            day_cell.fill = row_fill
            for r in range(start_row_for_day, current_row):
                ws_class.cell(row=r, column=1).border = thin_border

        # -------------------------------------------------------------
        # SHEET 2: THỜI KHÓA BIỂU THEO GIÁO VIÊN
        # -------------------------------------------------------------
        ws_teacher = wb.create_sheet(title="TKB_Theo_Giao_Vien")
        ws_teacher.views.sheetView[0].showGridLines = True

        ws_teacher.cell(row=1, column=1, value="BẢNG THỜI KHÓA BIỂU THEO GIÁO VIÊN").font = Font(size=14, bold=True, color="1F4E79")
        ws_teacher.cell(row=2, column=1, value="Sáng 4 tiết, Chiều 3 tiết - Chiều Thứ 6 nghỉ").font = Font(size=10, italic=True)

        t_row = 4
        ws_teacher.cell(row=t_row, column=1, value="Ngày").fill = header_fill
        ws_teacher.cell(row=t_row, column=1).font = header_font
        ws_teacher.cell(row=t_row, column=1).alignment = Alignment(horizontal='center', vertical='center')

        ws_teacher.cell(row=t_row, column=2, value="Buổi").fill = header_fill
        ws_teacher.cell(row=t_row, column=2).font = header_font
        ws_teacher.cell(row=t_row, column=2).alignment = Alignment(horizontal='center', vertical='center')

        ws_teacher.cell(row=t_row, column=3, value="Tiết").fill = header_fill
        ws_teacher.cell(row=t_row, column=3).font = header_font
        ws_teacher.cell(row=t_row, column=3).alignment = Alignment(horizontal='center', vertical='center')

        col_idx = 4
        for t in self.teachers:
            cell = ws_teacher.cell(row=t_row, column=col_idx, value=f"{t.name}")
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            col_idx += 1

        t_row += 1
        day_idx = 0
        for d in self.config.days:
            start_r = t_row
            row_fill = day_fill_toggle[day_idx % 2]
            day_idx += 1

            for session_name, periods in [("Sáng", self.config.morning_periods), ("Chiều", self.config.afternoon_periods)]:
                start_sess_r = t_row
                for p in periods:
                    ws_teacher.cell(row=t_row, column=3, value=f"Tiết {p}").alignment = Alignment(horizontal='center', vertical='center')
                    ws_teacher.cell(row=t_row, column=3).border = thin_border
                    ws_teacher.cell(row=t_row, column=3).fill = row_fill

                    col = 4
                    is_closed = (d, p) in self.config.closed_slots
                    for t in self.teachers:
                        cell = ws_teacher.cell(row=t_row, column=col)
                        cell.border = thin_border
                        cell.fill = row_fill
                        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

                        if is_closed:
                            cell.value = "[NGHỈ CHIỀU]"
                            cell.fill = closed_fill
                            cell.font = Font(color="7F7F7F", italic=True)
                        elif (d, p) in t.unavailable_slots:
                            cell.value = "[BẬN]"
                            cell.fill = PatternFill(start_color="FCE4D6", fill_type="solid")
                            cell.font = Font(color="C00000", italic=True)
                        else:
                            slots = teacher_slot_map.get((t.id, d, p), [])
                            if slots:
                                cell.value = "\n".join(f"{s.class_id} - {s.subject}" for s in slots)
                            else:
                                cell.value = ""
                        col += 1
                    t_row += 1

                ws_teacher.merge_cells(start_row=start_sess_r, start_column=2, end_row=t_row - 1, end_column=2)
                sess_cell = ws_teacher.cell(row=start_sess_r, column=2, value=session_name)
                sess_cell.alignment = Alignment(horizontal='center', vertical='center')
                sess_cell.font = Font(bold=True)
                sess_cell.fill = session_fill
                for r in range(start_sess_r, t_row):
                    ws_teacher.cell(row=r, column=2).border = thin_border

            ws_teacher.merge_cells(start_row=start_r, start_column=1, end_row=t_row - 1, end_column=1)
            day_cell = ws_teacher.cell(row=start_r, column=1, value=d)
            day_cell.alignment = Alignment(horizontal='center', vertical='center')
            day_cell.font = Font(bold=True)
            day_cell.fill = row_fill
            for r in range(start_r, t_row):
                ws_teacher.cell(row=r, column=1).border = thin_border

        # Căn chỉnh độ rộng cột
        for ws in [ws_class, ws_teacher]:
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        lines = str(cell.value).split("\n")
                        line_lens = [len(l) for l in lines]
                        max_len = max(max_len, max(line_lens))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(file_path)
        print(f"-> Đã xuất kết quả thành công ra file Excel: {file_path}")
