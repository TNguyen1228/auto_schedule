import os
from typing import Tuple, List, Set, Dict, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from scheduler.models import (
    ScheduleConfig,
    ClassGroup,
    Teacher,
    TeachingAssignment
)
from data.assignment_generator import (
    generate_phan_cong_df,
    generate_assignments_from_df
)


def parse_slot_str(slot_str: str) -> List[Tuple[str, int]]:
    """Phân tích chuỗi 'Thứ 2-T1, Thứ 6-T4' thành [('Thứ 2', 1), ('Thứ 6', 4)]"""
    if not slot_str or pd.isna(slot_str):
        return []
    slots = []
    parts = str(slot_str).split(",")
    for p in parts:
        p = p.strip()
        if "-" in p:
            day_part, period_part = p.split("-")
            day_part = day_part.strip()
            period_part = period_part.strip().replace("T", "").replace("Tiết", "").strip()
            try:
                slots.append((day_part, int(period_part)))
            except ValueError:
                pass
    return slots


def get_default_source_tables() -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """
    Trả về bộ 4 DataFrame nguồn chuẩn thực tế của trường:
    1. df_ctk (Bảng A: Chuong_Trinh_Khung)
    2. df_lop (Bảng B: Lop_Hoc - có GV Chủ Nhiệm)
    3. df_gvbm (Bảng C: Giao_Vien_Bo_Mon - map liên khối, bộ môn & ngoại lệ)
    4. df_gv (Danh sách Giáo viên & Lịch bận)
    """
    # Bảng B: Lop_Hoc
    df_lop = pd.DataFrame([
        {"Mã Lớp": "1A", "Tên Lớp": "Lớp 1A", "Khối": 1, "GV Chủ Nhiệm": "GV_1A"},
        {"Mã Lớp": "1B", "Tên Lớp": "Lớp 1B", "Khối": 1, "GV Chủ Nhiệm": "GV_1B"},
        {"Mã Lớp": "2A", "Tên Lớp": "Lớp 2A", "Khối": 2, "GV Chủ Nhiệm": "GV_2A"},
        {"Mã Lớp": "2B", "Tên Lớp": "Lớp 2B", "Khối": 2, "GV Chủ Nhiệm": "GV_2B"},
        {"Mã Lớp": "3A", "Tên Lớp": "Lớp 3A", "Khối": 3, "GV Chủ Nhiệm": "GV_3A"},
        {"Mã Lớp": "4A", "Tên Lớp": "Lớp 4A", "Khối": 4, "GV Chủ Nhiệm": "GV_4A"},
        {"Mã Lớp": "4B", "Tên Lớp": "Lớp 4B", "Khối": 4, "GV Chủ Nhiệm": "GV_4B"},
        {"Mã Lớp": "5A", "Tên Lớp": "Lớp 5A", "Khối": 5, "GV Chủ Nhiệm": "GV_5A"},
        {"Mã Lớp": "5B", "Tên Lớp": "Lớp 5B", "Khối": 5, "GV Chủ Nhiệm": "GV_5B"},
    ])

    # Bảng A: Chuong_Trinh_Khung
    df_ctk = pd.DataFrame([
        # --- KHỐI 1 (14 phân công = 32 tiết/tuần) ---
        {"Khối": "1", "Môn Học": "Chào cờ", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "GV_CHUNG", "Tiết Cố Định": "Thứ 2-T1", "Chung Toàn Trường": "Có"},
        {"Khối": "1", "Môn Học": "Sinh hoạt lớp", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "Thứ 6-T4", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "Toán", "Số Tiết/Tuần": 3, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "Tiếng Việt", "Số Tiết/Tuần": 12, "Tiết Tối Đa/Ngày": 2, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "HĐTN", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "Tiếng Anh", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:ANH", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "Âm nhạc", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:NHAC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "Mỹ thuật", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:HOA", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "TC Tiếng Việt", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "TC Toán", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 2, "Loại GV": "BO_MON:TC_TOAN", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "Đạo đức", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:DAO_DUC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "TC Tiếng Việt", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:TC_TV", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "GDTC", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:GDTC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "1", "Môn Học": "Tự nhiên & Xã hội", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 2, "Loại GV": "BO_MON:TNXH", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},

        # --- KHỐI 2 (13 phân công = 32 tiết/tuần) ---
        {"Khối": "2", "Môn Học": "Chào cờ", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "GV_CHUNG", "Tiết Cố Định": "Thứ 2-T1", "Chung Toàn Trường": "Có"},
        {"Khối": "2", "Môn Học": "Sinh hoạt lớp", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "Thứ 6-T4", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "Toán", "Số Tiết/Tuần": 5, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "Tiếng Việt", "Số Tiết/Tuần": 10, "Tiết Tối Đa/Ngày": 2, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "HĐTN", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "Tiếng Anh", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:ANH", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "Âm nhạc", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:NHAC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "Mỹ thuật", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:HOA", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "TC Toán", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 2, "Loại GV": "BO_MON:TC_TOAN", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "Đạo đức", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:DAO_DUC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "TC Tiếng Việt", "Số Tiết/Tuần": 3, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:TC_TV", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "GDTC", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:GDTC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "2", "Môn Học": "Tự nhiên & Xã hội", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 2, "Loại GV": "BO_MON:TNXH", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},

        # --- KHỐI 3 (16 phân công = 32 tiết/tuần) ---
        {"Khối": "3", "Môn Học": "Chào cờ", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "GV_CHUNG", "Tiết Cố Định": "Thứ 2-T1", "Chung Toàn Trường": "Có"},
        {"Khối": "3", "Môn Học": "Sinh hoạt lớp", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "Thứ 6-T4", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "Toán", "Số Tiết/Tuần": 5, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "Tiếng Việt", "Số Tiết/Tuần": 7, "Tiết Tối Đa/Ngày": 2, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "TTCM", "Số Tiết/Tuần": 3, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "HĐTN", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "Tiếng Anh", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:ANH", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "Âm nhạc", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:NHAC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "Mỹ thuật", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:HOA", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "Tin học", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:TIN", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "TC Tiếng Việt", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 2, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "TC Toán", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 2, "Loại GV": "BO_MON:TC_TOAN", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "TC Toán", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 2, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "GDTC", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:GDTC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "Công nghệ", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:CONG_NGHE", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "3", "Môn Học": "Tự nhiên & Xã hội", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:TNXH", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},

        # --- KHỐI 4 & 5 (16 phân công = 32 tiết/tuần) ---
        {"Khối": "4,5", "Môn Học": "Chào cờ", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "GV_CHUNG", "Tiết Cố Định": "Thứ 2-T1", "Chung Toàn Trường": "Có"},
        {"Khối": "4,5", "Môn Học": "Sinh hoạt lớp", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "Thứ 6-T4", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "Toán", "Số Tiết/Tuần": 5, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "Tiếng Việt", "Số Tiết/Tuần": 7, "Tiết Tối Đa/Ngày": 2, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "HĐTN", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "Tiếng Anh", "Số Tiết/Tuần": 4, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:ANH", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "Âm nhạc", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:NHAC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "Mỹ thuật", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:HOA", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "Tin học", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:TIN", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "TC Tiếng Việt", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "TC Toán", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "Đạo đức", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "CHU_NHIEM", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "GDTC", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:GDTC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "Công nghệ", "Số Tiết/Tuần": 1, "Tiết Tối Đa/Ngày": 1, "Loại GV": "BO_MON:CONG_NGHE", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "Khoa học", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 2, "Loại GV": "BO_MON:KHOA_HOC", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
        {"Khối": "4,5", "Môn Học": "LS & ĐL", "Số Tiết/Tuần": 2, "Tiết Tối Đa/Ngày": 2, "Loại GV": "BO_MON:LS_DL", "Tiết Cố Định": "", "Chung Toàn Trường": "Không"},
    ])

    # Bảng C: Giao_Vien_Bo_Mon
    df_gvbm = pd.DataFrame([
        # Môn chung toàn trường
        {"Khối/Lớp": "Tất cả", "Môn/Loại": "ANH", "Mã GV": "GV_ANH", "Ghi Chú": "Tiếng Anh toàn trường"},
        {"Khối/Lớp": "Tất cả", "Môn/Loại": "NHAC", "Mã GV": "GV_NHAC", "Ghi Chú": "Âm nhạc toàn trường"},
        {"Khối/Lớp": "Tất cả", "Môn/Loại": "HOA", "Mã GV": "GV_HOA", "Ghi Chú": "Mỹ thuật toàn trường"},
        {"Khối/Lớp": "Tất cả", "Môn/Loại": "TIN", "Mã GV": "GV_TIN", "Ghi Chú": "Tin học (Khối 3, 4, 5)"},

        # Thể dục (GDTC)
        {"Khối/Lớp": "1", "Môn/Loại": "GDTC", "Mã GV": "GV_12", "Ghi Chú": "GDTC Khối 1 do GV_12 dạy"},
        {"Khối/Lớp": "2", "Môn/Loại": "GDTC", "Mã GV": "PHT", "Ghi Chú": "GDTC Khối 2 do Phó Hiệu Trưởng dạy"},
        {"Khối/Lớp": "3,4,5", "Môn/Loại": "GDTC", "Mã GV": "GV_TD", "Ghi Chú": "GDTC Khối 3-5 do GV_TD chuyên trách"},

        # GV liên khối 1-2 (GV_12)
        {"Khối/Lớp": "1", "Môn/Loại": "TC_TOAN", "Mã GV": "GV_12", "Ghi Chú": "TC Toán K1"},
        {"Khối/Lớp": "1", "Môn/Loại": "TC_TV", "Mã GV": "GV_12", "Ghi Chú": "TC Tiếng Việt K1 (2 tiết)"},
        {"Khối/Lớp": "1", "Môn/Loại": "DAO_DUC", "Mã GV": "GV_12", "Ghi Chú": "Đạo đức K1"},
        {"Khối/Lớp": "1,2", "Môn/Loại": "TNXH", "Mã GV": "GV_12", "Ghi Chú": "TNXH K1, K2"},

        # GV liên khối 2-3 (GV_23)
        {"Khối/Lớp": "2", "Môn/Loại": "TC_TOAN", "Mã GV": "GV_23", "Ghi Chú": "TC Toán K2"},
        {"Khối/Lớp": "2", "Môn/Loại": "TC_TV", "Mã GV": "GV_23", "Ghi Chú": "TC Tiếng Việt K2"},
        {"Khối/Lớp": "2", "Môn/Loại": "DAO_DUC", "Mã GV": "GV_12", "Ghi Chú": "Đạo đức K2 mặc định GV_12"},
        {"Khối/Lớp": "2B", "Môn/Loại": "DAO_DUC", "Mã GV": "GV_23", "Ghi Chú": "NGOẠI LỆ: 2B học Đạo đức với GV_23"},
        {"Khối/Lớp": "3", "Môn/Loại": "TC_TOAN", "Mã GV": "GV_23", "Ghi Chú": "TC Toán K3 (2 tiết liên khối)"},
        {"Khối/Lớp": "3", "Môn/Loại": "CONG_NGHE", "Mã GV": "GV_23", "Ghi Chú": "Công nghệ K3"},
        {"Khối/Lớp": "3", "Môn/Loại": "TNXH", "Mã GV": "GV_23", "Ghi Chú": "TNXH K3"},

        # GV liên khối 4-5 (GV_45)
        {"Khối/Lớp": "4,5", "Môn/Loại": "CONG_NGHE", "Mã GV": "GV_45", "Ghi Chú": "Công nghệ K4, K5"},
        {"Khối/Lớp": "4,5", "Môn/Loại": "KHOA_HOC", "Mã GV": "GV_45", "Ghi Chú": "Khoa học K4, K5"},
        {"Khối/Lớp": "4,5", "Môn/Loại": "LS_DL", "Mã GV": "GV_45", "Ghi Chú": "Lịch sử & Địa lý K4, K5"},
    ])

    # Danh sách Giáo viên & Lịch bận
    df_gv = pd.DataFrame([
        {"Mã GV": "PHT", "Tên Giáo Viên": "Hà", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_1A", "Tên Giáo Viên": "Đỗ Hương", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_1B", "Tên Giáo Viên": "Hà Hương", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_2A", "Tên Giáo Viên": "Thúy", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_2B", "Tên Giáo Viên": "Chiên", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_3A", "Tên Giáo Viên": "Nhiên", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_4A", "Tên Giáo Viên": "Lan", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_4B", "Tên Giáo Viên": "Nhật", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_5A", "Tên Giáo Viên": "Thanh", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_5B", "Tên Giáo Viên": "Vũ Hương", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_12", "Tên Giáo Viên": "Phượng", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_23", "Tên Giáo Viên": "Lập", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_45", "Tên Giáo Viên": "Hiền", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_ANH", "Tên Giáo Viên": "Nõn", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_TD", "Tên Giáo Viên": "Tấm", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_NHAC", "Tên Giáo Viên": "Tươi", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_HOA", "Tên Giáo Viên": "Việt", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_TIN", "Tên Giáo Viên": "Nga", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
        {"Mã GV": "GV_CHUNG", "Tên Giáo Viên": "BGH", "Tiết Bận Đăng Ký (VD: Thứ 5-T3, Thứ 5-T4)": ""},
    ])

    return df_ctk, df_lop, df_gvbm, df_gv


def create_excel_input_template(file_path: str = "Input_Mau_TKB.xlsx"):
    """
    Tạo file Excel mẫu chuẩn theo cấu trúc 3 BẢNG NGUỒN + TỰ SINH PHAN_CONG:
    - Sheet 1: Chuong_Trinh_Khung (Bảng A)
    - Sheet 2: Lop_Hoc (Bảng B - có cột GV Chủ Nhiệm)
    - Sheet 3: Giao_Vien_Bo_Mon (Bảng C - map loại GV bộ môn/liên khối và ngoại lệ lớp)
    - Sheet 4: Giao_Vien (Danh sách GV và lịch bận)
    - Sheet 5: Phan_Cong (Sheet tự sinh tự động từ A + B + C để kiểm tra và solver đọc)
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    df_ctk, df_lop, df_gvbm, df_gv = get_default_source_tables()

    # Sinh trước DataFrame Phan_Cong tự động từ 3 bảng A, B, C
    df_pc = generate_phan_cong_df(df_ctk, df_lop, df_gvbm)

    sheets_config = [
        ("Chuong_Trinh_Khung", df_ctk),
        ("Lop_Hoc", df_lop),
        ("Giao_Vien_Bo_Mon", df_gvbm),
        ("Giao_Vien", df_gv),
        ("Phan_Cong", df_pc),
    ]

    for title, df in sheets_config:
        ws = wb.create_sheet(title=title)
        ws.views.sheetView[0].showGridLines = True

        # Ghi header
        headers = list(df.columns)
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Ghi dữ liệu
        for row_idx, row in enumerate(df.itertuples(index=False), 2):
            for col_idx, val in enumerate(row, 1):
                val_to_write = "" if pd.isna(val) else val
                cell = ws.cell(row=row_idx, column=col_idx, value=val_to_write)
                cell.border = thin_border
                
                # Căn giữa các cột mã, số, loại, cố định
                if isinstance(val_to_write, (int, float)) or col_idx in (1, 3, 4, 5, 6, 7):
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

        # Căn chỉnh độ rộng cột
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(file_path)
    print(f"-> Đã tạo file Excel mẫu cấu trúc 3 bảng thành công: {file_path}")
    print(f"   + Bảng A (Chuong_Trinh_Khung): {len(df_ctk)} dòng khung chương trình")
    print(f"   + Bảng B (Lop_Hoc): {len(df_lop)} lớp kèm GV Chủ nhiệm")
    print(f"   + Bảng C (Giao_Vien_Bo_Mon): {len(df_gvbm)} dòng ánh xạ GV bộ môn/liên khối")
    print(f"   + Danh sách GV (Giao_Vien): {len(df_gv)} giáo viên")
    print(f"   + Tự sinh tự động (Phan_Cong): {len(df_pc)} phân công chi tiết (100% đầy đủ)")


def sync_phan_cong_to_excel(file_path: str = "Input_Mau_TKB.xlsx") -> int:
    """
    Đọc 3 bảng nguồn (Chuong_Trinh_Khung, Lop_Hoc, Giao_Vien_Bo_Mon) từ file_path,
    tự sinh ra sheet 'Phan_Cong' và ghi đè sheet 'Phan_Cong' trong file Excel đó.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Không tìm thấy file: {file_path}")

    xls = pd.ExcelFile(file_path)
    if "Chuong_Trinh_Khung" not in xls.sheet_names:
        raise ValueError(f"File '{file_path}' không có sheet 'Chuong_Trinh_Khung'.")

    df_ctk = pd.read_excel(file_path, sheet_name="Chuong_Trinh_Khung")
    df_lop = pd.read_excel(file_path, sheet_name="Lop_Hoc")
    df_gvbm = pd.read_excel(file_path, sheet_name="Giao_Vien_Bo_Mon")

    # Sinh Phan_Cong
    df_pc = generate_phan_cong_df(df_ctk, df_lop, df_gvbm)

    # Ghi đè sheet Phan_Cong vào file Excel giữ nguyên các sheet khác
    wb = openpyxl.load_workbook(file_path)
    if "Phan_Cong" in wb.sheetnames:
        del wb["Phan_Cong"]

    ws = wb.create_sheet(title="Phan_Cong")
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = list(df_pc.columns)
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')

    for row_idx, row in enumerate(df_pc.itertuples(index=False), 2):
        for col_idx, val in enumerate(row, 1):
            val_to_write = "" if pd.isna(val) else val
            cell = ws.cell(row=row_idx, column=col_idx, value=val_to_write)
            cell.border = thin_border
            if isinstance(val_to_write, (int, float)) or col_idx in (1, 3, 4, 5, 6, 7):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(file_path)
    print(f"-> Đã đồng bộ thành công {len(df_pc)} dòng vào sheet 'Phan_Cong' của file: {file_path}")
    return len(df_pc)


def load_data_from_excel(file_path: str) -> Tuple[ScheduleConfig, List[ClassGroup], List[Teacher], List[TeachingAssignment]]:
    """
    Đọc dữ liệu từ file Excel:
    - Nếu có sheet 'Chuong_Trinh_Khung': tự động sinh phân công từ 3 bảng nguồn A, B, C.
    - Nếu không có: đọc trực tiếp từ sheet 'Phan_Cong' (tương thích ngược hoàn hảo).
    """
    config = ScheduleConfig(
        days=['Thứ 2', 'Thứ 3', 'Thứ 4', 'Thứ 5', 'Thứ 6'],
        morning_periods=[1, 2, 3, 4],
        afternoon_periods=[5, 6, 7],
        closed_slots={('Thứ 6', 5), ('Thứ 6', 6), ('Thứ 6', 7)}
    )

    xls = pd.ExcelFile(file_path)

    # 1. Đọc Lop_Hoc
    df_classes = pd.read_excel(file_path, sheet_name="Lop_Hoc")
    classes = []
    col_grade = [c for c in df_classes.columns if "Khối" in c][0]
    col_cn = [c for c in df_classes.columns if "Chủ Nhiệm" in c]
    for _, row in df_classes.iterrows():
        cn_id = str(row[col_cn[0]]).strip() if col_cn and pd.notna(row[col_cn[0]]) else None
        classes.append(ClassGroup(
            id=str(row["Mã Lớp"]).strip(),
            name=str(row["Tên Lớp"]).strip(),
            grade=int(row[col_grade]),
            homeroom_teacher_id=cn_id
        ))

    # 2. Đọc Giao_Vien
    df_teachers = pd.read_excel(file_path, sheet_name="Giao_Vien")
    teachers = []
    for _, row in df_teachers.iterrows():
        busy_col = [c for c in df_teachers.columns if "Bận" in c or "Nghỉ" in c]
        busy_val = row[busy_col[0]] if busy_col else ""
        busy_slots = set(parse_slot_str(busy_val))
        teachers.append(Teacher(
            id=str(row["Mã GV"]).strip(),
            name=str(row["Tên Giáo Viên"]).strip(),
            unavailable_slots=busy_slots
        ))

    # 3. Phân công giảng dạy
    if "Chuong_Trinh_Khung" in xls.sheet_names and "Giao_Vien_Bo_Mon" in xls.sheet_names:
        # CẤU TRÚC MỚI: Tự sinh từ 3 bảng A, B, C
        df_ctk = pd.read_excel(file_path, sheet_name="Chuong_Trinh_Khung")
        df_gvbm = pd.read_excel(file_path, sheet_name="Giao_Vien_Bo_Mon")
        df_pc = generate_phan_cong_df(df_ctk, df_classes, df_gvbm)
        assignments = generate_assignments_from_df(df_pc)
    else:
        # CẤU TRÚC CŨ: Đọc trực tiếp từ sheet Phan_Cong
        df_assign = pd.read_excel(file_path, sheet_name="Phan_Cong")
        assignments = []
        for idx, row in df_assign.iterrows():
            fixed_col = [c for c in df_assign.columns if "Cố Định" in c]
            fixed_val = row[fixed_col[0]] if fixed_col else ""
            fixed = parse_slot_str(fixed_val)
            
            shared_col = [c for c in df_assign.columns if "Chung" in c]
            shared_val = row[shared_col[0]] if shared_col else "Không"
            shared = str(shared_val).strip().lower() in ("có", "co", "yes", "true", "1")
            
            max_col = [c for c in df_assign.columns if "Tối Đa" in c]
            max_p = int(row[max_col[0]]) if max_col and not pd.isna(row[max_col[0]]) else 2
            
            assignments.append(TeachingAssignment(
                id=f"AS_{idx+1:03d}",
                class_id=str(row["Mã Lớp"]).strip(),
                subject=str(row["Môn Học"]).strip(),
                teacher_id=str(row["Mã GV"]).strip(),
                periods_per_week=int(row["Số Tiết/Tuần"]),
                max_periods_per_day=max_p,
                fixed_slots=fixed,
                is_shared_activity=shared
            ))

    return config, classes, teachers, assignments
